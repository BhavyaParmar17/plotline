"""CSV and Excel exports of the calendar for review."""
from __future__ import annotations

import csv
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .rules import THEMES, WEEKDAYS

COLUMNS = [
    ("Date", "date", 12), ("Day", "weekday", 6), ("Time", "time", 7), ("Platform", "platform_label", 15),
    ("Format", "format_label", 17), ("Pillar", "theme_label", 18), ("Funnel stage", "funnel", 14),
    ("Campaign", "campaign", 20), ("Occasion", "occasion", 22), ("Service", "service", 22), ("Idea", "title", 44),
    ("Brief", "brief", 60), ("How to make it", "how_to", 50), ("Call to action", "cta", 30),
    ("Asset", "asset", 22), ("Asset to source", "asset_gap", 14), ("Status", "status", 10), ("Flags", "flags", 40),
]


def _flag_map(analysis: dict) -> dict:
    out = {}
    for f in analysis["flags"]:
        for pid in f["posts"]:
            out.setdefault(pid, []).append(f["message"])
    return out


def _value(post, key, flags):
    if key == "flags":
        return " | ".join(flags.get(post["id"], []))
    if key == "asset_gap":
        return "Yes" if post.get("asset_gap") else ""
    return post.get(key, "") or ""


def to_csv(posts: list[dict], analysis: dict) -> str:
    flags = _flag_map(analysis)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c[0] for c in COLUMNS])
    for p in posts:
        w.writerow([_value(p, c[1], flags) for c in COLUMNS])
    return "\ufeff" + buf.getvalue()


def _hex(color: str) -> str:
    return color.lstrip("#").upper()


def _tint(color: str, amount: float = 0.82) -> str:
    c = color.lstrip("#")
    r, g, b = (int(c[i:i + 2], 16) for i in (0, 2, 4))
    mix = lambda v: int(v + (255 - v) * amount)  # noqa: E731
    return f"{mix(r):02X}{mix(g):02X}{mix(b):02X}"


def to_xlsx(brief: dict, posts: list[dict], analysis: dict) -> bytes:
    flags = _flag_map(analysis)
    wb = Workbook()
    thin = Side(style="thin", color="CBD6CF")
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="142019")

    # Calendar sheet
    ws = wb.active
    ws.title = "Calendar"
    ws.append([c[0] for c in COLUMNS])
    for p in posts:
        ws.append([_value(p, c[1], flags) for c in COLUMNS])
    for i, (_, _, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for cell in ws[1]:
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center")
    theme_col = [c[1] for c in COLUMNS].index("theme_label") + 1
    for r, p in enumerate(posts, start=2):
        for c in range(1, len(COLUMNS) + 1):
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
        color = THEMES[p["theme"]]["color"]
        ws.cell(r, theme_col).fill = PatternFill("solid", fgColor=_tint(color))
        ws.cell(r, theme_col).font = Font(bold=True, color=_hex(color))
        if flags.get(p["id"]):
            ws.cell(r, len(COLUMNS)).font = Font(color="B4233C")
    if posts:
        ref = f"A1:{get_column_letter(len(COLUMNS))}{len(posts) + 1}"
        table = Table(displayName="Calendar", ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
        ws.add_table(table)
    ws.freeze_panes = "B2"

    # Month view sheet
    mv = wb.create_sheet("Month view")
    first = date.fromisoformat(brief["first_day"])
    last = date.fromisoformat(brief["last_day"])
    mv.append([f"{brief['business']} — {first.strftime('%B %Y')}"])
    mv["A1"].font = Font(bold=True, size=14)
    mv.append(WEEKDAYS)
    for cell in mv[2]:
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center")
    by_day = {}
    for p in posts:
        by_day.setdefault(p["date"], []).append(p)
    row = 3
    col = first.weekday() + 1
    d = first
    while d <= last:
        items = by_day.get(d.isoformat(), [])
        lines = [str(d.day)] + [f"{p['platform_label']}, {p['format_label']}: {p['title']}" for p in items]
        cell = mv.cell(row, col, "\n".join(lines))
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        if items:
            cell.fill = PatternFill("solid", fgColor=_tint(THEMES[items[0]["theme"]]["color"], 0.88))
        mv.row_dimensions[row].height = max(mv.row_dimensions[row].height or 0, 30 + 42 * len(items))
        col += 1
        if col > 7:
            col = 1
            row += 1
        d = date.fromordinal(d.toordinal() + 1)
    for i in range(1, 8):
        mv.column_dimensions[get_column_letter(i)].width = 30

    # Balance sheet
    bs = wb.create_sheet("Balance")
    bs.append(["Balance score", analysis["score"], analysis["score_label"]])
    bs["A1"].font = Font(bold=True)
    bs.append([])
    bs.append(["Pillar", "Posts", "Share %", "Target %", "Difference"])
    for cell in bs[3]:
        cell.font = head_font
        cell.fill = head_fill
    for t in analysis["themes"]:
        bs.append([t["label"], t["count"], t["pct"], t["target_pct"], t["diff"]])
    bs.append([])
    bs.append(["Platform", "Format", "Posts", "Share %", "Target %"])
    for cell in bs[bs.max_row]:
        cell.font = head_font
        cell.fill = head_fill
    for f in analysis["formats"]:
        for row_ in f["formats"]:
            bs.append([f["label"], row_["label"], row_["count"], row_["pct"], row_["target_pct"]])
    bs.append([])
    bs.append(["Service", "Posts"])
    for cell in bs[bs.max_row]:
        cell.font = head_font
        cell.fill = head_fill
    for s in analysis["services"]:
        bs.append([s["name"], s["count"]])
    for i, w in enumerate([24, 22, 10, 10, 12], start=1):
        bs.column_dimensions[get_column_letter(i)].width = w

    # Flags sheet
    fs = wb.create_sheet("Flags")
    fs.append(["Severity", "Type", "Posts", "What’s wrong", "Suggestion"])
    for cell in fs[1]:
        cell.font = head_font
        cell.fill = head_fill
    for f in analysis["flags"]:
        fs.append([f["severity"], f["type"], ", ".join(f["posts"]), f["message"], f["suggestion"]])
    for i, w in enumerate([10, 12, 24, 70, 50], start=1):
        fs.column_dimensions[get_column_letter(i)].width = w
    for r in fs.iter_rows(min_row=2):
        for cell in r:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
